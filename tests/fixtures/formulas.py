from pathlib import Path
from openpyxl import Workbook
root=Path(__file__).parent
wb=Workbook();ws=wb.active;ws.title='Cases'
for i,value in enumerate([1,2,'3',True,None],2):ws.cell(i,1,value);ws.cell(i,2,i*10)
formulas=['SUM(A2:A6)','SUM(A4)','SUM(A5)','SUM("3",TRUE)','COUNT(A4)','COUNT(A5)','COUNT("3",TRUE,"bad")','COUNTA(A2:A6)','AVERAGE(A2:A6)','MIN(A9:A10)','MAX(A9:A10)','MIN(A2:A6)','MAX(A2:A6)','IF(TRUE,10,1/0)','IF(FALSE,1/0,20)','IFERROR(1/0,99)','_xlfn.IFNA(VLOOKUP("missing",A2:B6,2,FALSE),0)','AND(A2:A6)','OR(A2:A6)','NOT(FALSE)','ABS(-3)','ROUND(1.005,2)','ROUND(-1.5,0)','ROUND(1234,-2)','LEN("abc")','LOWER("ABC")','UPPER("abc")','TRIM("  a   b  ")','_xlfn.CONCAT("a","b")','CONCATENATE("a","b")','COUNTIF(A2:A6,">0")','COUNTIF(A2:A6,3)','SUMIF(A2:A6,">0",B2:B6)','INDEX(A2:B6,2,2)','MATCH(2,A2:A6,0)','VLOOKUP(2,A2:B6,2,FALSE)','HLOOKUP(1,A2:B6,2,FALSE)','_xlfn.XLOOKUP(2,A2:A6,B2:B6)','_xlfn.XLOOKUP(99,A2:A6,B2:B6,"missing")','A9=0','A9=""','2^3^2','20%*50','"a"&"b"','"abc"="ABC"','A9','A2+A3*2','(A2+A3)*2','SUM(A2:A6,4)','AVERAGE(A4,4)','COUNT(A2:A6)','IFERROR(MATCH(99,A2:A6,0),0)']
for i,f in enumerate(formulas,1):ws.cell(i,3,'='+f)
wb.save(root/'formula-oracle-input.xlsx')
