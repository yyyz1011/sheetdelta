"""Generate project-owned fixtures with openpyxl 3.1.5; no third-party user data."""
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from zipfile import ZipFile, ZIP_DEFLATED
import re
root=Path(__file__).parent
wb=Workbook(); ws=wb.active; ws.title='Data'
ws.append(['id','amount','date','active','note'])
ws.append(['001',12.5,datetime(2024,2,29,12,30),True,'line\n& <>'])
ws['C2'].number_format='yyyy-mm-dd hh:mm'
ws.append([None]*5)
ws.append(['900719925474099312345',0,None,False,'text'])
ws=wb.create_sheet('Hidden'); ws.append(['id','value']); ws.append(['h',1]); ws.sheet_state='hidden'
ws=wb.create_sheet('Merged'); ws.append(['id','group']); ws.append(['1','Group']); ws.append(['2',None]); ws.merge_cells('B2:B3')
ws=wb.create_sheet('Formulas'); ws.append(['id','cached','missing']); ws.append(['1','=SUM(2,3)','=TODAY()'])
ws=wb.create_sheet('Errors'); ws.append(['id','value']); ws.append(['1','#DIV/0!'])
ws=wb.create_sheet('Titles'); ws.append(['Report',None]); ws.merge_cells('A1:B1'); ws.append(['id','value']); ws.append(['001',2])
wb.save(root/'openpyxl-compat.xlsx')
# openpyxl does not calculate formulas. Add a known cached result to one formula cell.
p=root/'openpyxl-compat.xlsx'
with ZipFile(p) as z: contents={n:z.read(n) for n in z.namelist()}
name='xl/worksheets/sheet4.xml';xml=contents[name].decode();xml=re.sub(r'(<c r="B2"[^>]*>.*?)(?:<v\s*/>|<v>.*?</v>)',r'\g<1><v>5</v>',xml);contents[name]=xml.encode()
with ZipFile(p,'w',ZIP_DEFLATED) as z:
 for name,data in contents.items(): z.writestr(name,data)
wb=Workbook();wb.epoch=CALENDAR_MAC_1904;ws=wb.active;ws.title='Dates';ws.append(['id','date']);ws.append(['001',datetime(2024,2,29)]);ws['B2'].number_format='yyyy-mm-dd';wb.save(root/'openpyxl-1904.xlsx')
(root/'gb18030.csv').write_bytes('编号,名称\r\n001,测试商品\r\n'.encode('gb18030'))
