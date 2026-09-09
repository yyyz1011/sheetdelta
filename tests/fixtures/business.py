"""Project-owned invoice workbook for cross-application roundtrips."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
root=Path(__file__).parent
wb=Workbook();ws=wb.active;ws.title='Orders'
ws.append(['sku','qty','unit_price','total','label'])
for row in [('001',2,12.5,'=B2*C2','=IF(D2>20,"large","small")'),('002',3,5,'=B3*C3','=IF(D3>20,"large","small")'),('003',0,7,'=B4*C4','=IF(D4>20,"large","small")')]:ws.append(row)
for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='2459CC')
for cell in ws['C'][1:]+ws['D'][1:]:cell.number_format='0.00'
ws['A2'].comment=Comment('Keep identifiers as text','SheetDelta')
ws['A2'].hyperlink='https://example.com/products/001'
ws.column_dimensions['A'].width=24;ws.freeze_panes='B2';ws.auto_filter.ref='A1:E4'
dv=DataValidation(type='whole',operator='greaterThanOrEqual',formula1=0);ws.add_data_validation(dv);dv.add('B2:B100')
chart=BarChart();chart.title='Order totals';chart.add_data(Reference(ws,min_col=4,min_row=1,max_row=4),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=4));ws.add_chart(chart,'G2')
summary=wb.create_sheet('Summary');summary.append(['metric','value']);summary.append(['revenue','=SUM(Orders!D2:D4)']);summary.append(['average','=AVERAGE(Orders!D2:D4)']);summary.append(['large','=COUNTIF(Orders!D2:D4,">20")']);summary.append(['rounded','=ROUND(B3,2)']);summary.append(['safe','=IFERROR(1/0,0)'])
hidden=wb.create_sheet('Internal');hidden.append(['id','value']);hidden.append(['rate',0.1]);hidden.sheet_state='hidden'
wb.save(root/'business-template.xlsx')
